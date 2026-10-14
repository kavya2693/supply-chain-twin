#!/usr/bin/env python3
"""Build the slide-by-slide audit report for a repo demo video.

Reads one per-repo facts file (YAML) and writes, into project.out_dir:
  slides/slide-XX.jpg      one still per slide (from the rendered master, if `stills` is given)
  AUDIT_REPORT.md          pastable Markdown (GitHub)
  AUDIT_REPORT.docx        Word version, images embedded
  AUDIT_REPORT.pdf         converted from the .docx with LibreOffice
  audit_checks.json        machine-readable check results

Every "direct"/"derived" figure with a `check` is recomputed from the repo outputs and compared
with what the slide shows; figures without a check are marked "manual" with the reason.

Usage:  python3 build_audit_report.py <facts.yaml> [--master renders/master.mp4] [--no-stills] [--no-pdf]
Relative paths in the facts file (repo_dir, demo_dir, out_dir) are resolved against the facts file's folder.
Needs (system python3): pyyaml, pandas, python-docx, Pillow; ffmpeg; LibreOffice (soffice).
"""
# ruff: noqa: E501, N802, N806  (R/J/T/P are the short helper names used by facts.yaml check expressions)

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pandas as pd
import yaml

FFMPEG = shutil.which("ffmpeg") or "/opt/homebrew/bin/ffmpeg"
SOFFICE = shutil.which("soffice") or "/opt/homebrew/bin/soffice"

TYPE_LABEL = {
    "direct": "direct",
    "derived": "derived",
    "illustrative": "illustrative",
    "general-knowledge": "general knowledge",
}


# ----------------------------------------------------------------------------- checks
def make_env(repo: Path, demo: Path, helpers: str | None) -> dict:
    def P(p: str) -> Path:
        return demo / p[5:] if p.startswith("demo:") else repo / p

    def R(p):
        return pd.read_csv(P(p))

    def J(p):
        return json.loads(P(p).read_text())

    def T(p):
        return P(p).read_text()

    def grep(p, rx):
        return re.search(rx, T(p), re.M) is not None

    def row(p, **eq):
        df = R(p)
        for k, v in eq.items():
            df = df[df[k] == v]
        if len(df) != 1:
            raise ValueError(f"{p}: {len(df)} rows match {eq}")
        return df.iloc[0]

    env = {"R": R, "J": J, "T": T, "grep": grep, "row": row, "pd": pd, "json": json, "re": re}
    if helpers:
        exec(helpers, env)
    return env


def run_checks(facts: dict, repo: Path, demo: Path) -> list[dict]:
    env = make_env(repo, demo, facts.get("helpers"))
    out = []
    for f in facts["figures"]:
        rec = {"id": f["id"], "shown": f["shown"], "type": f["type"]}
        chk = f.get("check")
        if not chk:
            rec.update(status="manual", computed="-", reason=f.get("manual", "no automatic source"))
        else:
            try:
                val = eval(chk["expr"], env)
                if chk["fmt"] == "bool":
                    ok = val is True or (hasattr(val, "item") and bool(val) is True)
                    rec.update(computed="true" if ok else "false", expect="true")
                else:
                    if "{:d}" in chk["fmt"] or "{:,d}" in chk["fmt"]:
                        val = int(val)
                    s = chk["fmt"].format(float(val) if "f}" in chk["fmt"] else val)
                    ok = s == chk["expect"]
                    rec.update(computed=s, expect=chk["expect"])
                rec["status"] = "pass" if ok else "FAIL"
            except Exception as e:
                rec.update(status="FAIL", computed=f"error: {e}")
        out.append(rec)
    return out


# ----------------------------------------------------------------------------- environment
def environment(facts: dict, repo: Path) -> dict:
    p = facts["project"]
    env = {}

    def sh(*cmd):
        try:
            return subprocess.run(cmd, cwd=repo, capture_output=True, text=True, check=True).stdout.strip()
        except Exception:
            return "n/a"

    # anchor on the commit that last changed the outputs: it is stable, unlike the working HEAD
    env["commit"] = sh("git", "log", "-1", "--format=%H", "--", p.get("outputs_dir", "outputs"))
    env["commit_short"] = env["commit"][:7]
    env["outputs_commit"] = sh("git", "log", "-1", "--format=%h (%ci)", "--", p.get("outputs_dir", "outputs"))
    # uncommitted files outside the report's own output folder (docs/demo is written by this script)
    env["dirty"] = "\n".join(
        ln for ln in sh("git", "status", "--short").splitlines() if "docs/demo" not in ln
    )
    py = repo / p.get("venv_python", ".venv/bin/python")
    env["python"] = sh(str(py), "-c", "import sys; print(sys.version.split()[0])")
    pk = p.get("version_packages", [])
    code = "import importlib\n" + "\n".join(
        f"m=importlib.import_module('{m}'); print('{m}', getattr(m,'__version__','?'))" for m in pk
    )
    env["packages"] = sh(str(py), "-c", code)
    return env


# ----------------------------------------------------------------------------- stills
def make_stills(facts: dict, demo: Path, out: Path) -> None:
    st = facts.get("stills")
    if not st:
        return
    master = demo / st["master"]
    (out / "slides").mkdir(parents=True, exist_ok=True)
    for key, t in st["times"].items():
        dst = out / "slides" / f"slide-{key}.jpg"
        subprocess.run(
            [
                FFMPEG,
                "-v",
                "error",
                "-y",
                "-ss",
                str(t),
                "-i",
                str(master),
                "-frames:v",
                "1",
                "-vf",
                f"scale={st.get('width', 1600)}:-1:flags=lanczos",
                "-q:v",
                "3",
                str(dst),
            ],
            check=True,
        )


# ----------------------------------------------------------------------------- helpers
def fig_by_id(facts):
    return {f["id"]: f for f in facts["figures"]}


def slides_of(ids, figs):
    s = sorted({sl for i in ids if i in figs for sl in figs[i]["slides"]})
    return s


def impact_rows(facts):
    figs = fig_by_id(facts)
    rows = []
    for a in facts["assumptions"]:
        ids = list(figs) if a.get("all_figures") else a.get("figures", [])
        slides = slides_of(ids, figs)
        rows.append(
            {
                "id": a["id"],
                "input": a["statement"].split(".")[0],
                "figures": "all figures" if a.get("all_figures") else ", ".join(ids),
                "slides": ", ".join(slides),
                "edit": "; ".join(a.get("edit", [])),
                "rerun": "; ".join(a.get("rerun", [])),
            }
        )
    return rows


def assumption_figs(a, figs):
    return "all figures" if a.get("all_figures") else ", ".join(a.get("figures", []))


def md_escape(s: str) -> str:
    return str(s).replace("|", "\\|").replace("\n", " ")


# ----------------------------------------------------------------------------- markdown
def write_md(facts, checks, env, out: Path, generated: str) -> Path:
    p = facts["project"]
    figs = fig_by_id(facts)
    chk = {c["id"]: c for c in checks}
    n_pass = sum(c["status"] == "pass" for c in checks)
    n_fail = sum(c["status"] == "FAIL" for c in checks)
    n_man = sum(c["status"] == "manual" for c in checks)
    L = []
    a = L.append
    a(f"# {p['title']}\n")
    a(f"**Repository:** {p['repo_url']}  ")
    a(
        f"**Numbers taken from commit:** `{env['commit_short']}` (outputs last changed in `{env['outputs_commit']}`)  "
    )
    a(f"**Video:** `{p['video']}` ({p['video_duration_s']:.1f} s, 1920x1080, 30 fps, silent)  ")
    a(f"**Generated:** {generated}  ")
    a(f"**Automated checks:** {n_pass} pass · {n_fail} fail · {n_man} manual (of {len(checks)} figures)\n")
    a("## How to use this report\n")
    a(
        "Every number or claim on screen has an ID (`F-01`, `F-02`, ...) and every modelling assumption has an ID "
        "(`A-01`, ...). Cite the ID. Each figure lists its source file, the exact column/row/key, the formula for "
        "derived numbers, and an automated check result. The change-impact matrix tells you which slides and files "
        "change if an input or assumption changes.\n"
    )
    a(
        "Types: **direct** = read straight from a repo output; **derived** = computed from direct values (formula "
        "given); **illustrative** = drawing or example, not data; **general knowledge** = context not produced by the repo.\n"
    )
    a("### Regenerate\n")
    a("```bash")
    for c in p.get("regenerate", []):
        a(c)
    a("```\n")

    a("## 1. Figures\n")
    a("| ID | Shown on screen | Slides | Source | Column / row / key | How computed | Type | Check |")
    a("|---|---|---|---|---|---|---|---|")
    for f in facts["figures"]:
        c = chk[f["id"]]
        status = {"pass": "✓ verified", "FAIL": "✗ FAIL", "manual": "manual"}[c["status"]]
        if c["status"] == "manual":
            status += f": {c['reason']}"
        a(
            f"| {f['id']} | {md_escape(f['shown'])} | {', '.join(f['slides'])} | `{md_escape(f['source'])}` | "
            f"{md_escape(f['key'])} | {md_escape(f['formula'])} | {TYPE_LABEL[f['type']]} | {md_escape(status)} |"
        )
    a("")

    a(f"## 2. Automated check run ({generated})\n")
    a("| ID | Expected on slide | Recomputed from repo | Result |")
    a("|---|---|---|---|")
    for c in checks:
        res = {"pass": "✓ pass", "FAIL": "✗ FAIL", "manual": "manual"}[c["status"]]
        a(f"| {c['id']} | {md_escape(c.get('expect', c['shown']))} | {md_escape(c['computed'])} | {res} |")
    a("")

    a("## 3. Assumptions register\n")
    a("| ID | Assumption | Source | Used by | Sensitivity: if this changes |")
    a("|---|---|---|---|---|")
    for x in facts["assumptions"]:
        ids = list(figs) if x.get("all_figures") else x.get("figures", [])
        a(
            f"| {x['id']} | {md_escape(x['statement'])} | `{md_escape(x['source'])}` | {assumption_figs(x, figs)} | "
            f"figures {assumption_figs(x, figs)} and slides {', '.join(slides_of(ids, figs))} change |"
        )
    a("")

    a("## 4. Slide by slide\n")
    for s in facts["slides"]:
        a(f"### Slide {s['id']}: {s['title']}  \n*{s['time']} · composition `{s['composition']}`*\n")
        a(f"![Slide {s['id']}]({s['image']})\n")
        if s.get("image2"):
            a(f"![Slide {s['id']} (second phase)]({s['image2']})\n")
        a("**On-screen text (verbatim)**\n")
        for t in s["text"]:
            a(f"> {t}  ")
        a("")
        on = [f for f in facts["figures"] if s["id"] in f["slides"]]
        if on:
            a("**Figures on this slide**\n")
            a("| ID | Shown | Source · key | Calculation | Check |")
            a("|---|---|---|---|---|")
            for f in on:
                c = chk[f["id"]]
                res = {"pass": "✓", "FAIL": "✗", "manual": "manual"}[c["status"]]
                a(
                    f"| {f['id']} | {md_escape(f['shown'])} | `{md_escape(f['source'])}` · {md_escape(f['key'])} | "
                    f"{md_escape(f['formula'])} | {res} |"
                )
            a("")
            notes = [f for f in on if f.get("note")]
            for f in notes:
                a(f"- *{f['id']} note:* {f['note']}")
            if notes:
                a("")
        a(f"**Assumptions:** {', '.join(s.get('assumptions', [])) or '-'}\n")
        if s.get("illustrative"):
            a("**Illustrative, not data:**\n")
            for t in s["illustrative"]:
                a(f"- {t}")
            a("")
        a(
            f"**To change this slide:** {s['change']} (project folder `{p.get('demo_display', p['demo_dir'])}`).\n"
        )

    a("## 5. Change-impact matrix\n")
    a("| Input / assumption | Affected figures | Affected slides | Files to edit | Commands to re-run |")
    a("|---|---|---|---|---|")
    for r in impact_rows(facts):
        a(
            f"| {r['id']}: {md_escape(r['input'])} | {r['figures']} | {r['slides']} | `{md_escape(r['edit'])}` | `{md_escape(r['rerun'])}` |"
        )
    a("")

    a("## 6. Known differences and caveats\n")
    for c in facts.get("caveats", []):
        a(f"- {c}")
    a("")

    a("## Appendix: environment\n")
    a(f"- Commit: `{env['commit']}`")
    if env["dirty"]:
        a("- Working tree has uncommitted files (demo assets): `" + env["dirty"].replace("\n", "; ") + "`")
    a(f"- Python: {env['python']}")
    for ln in env["packages"].splitlines():
        a(f"- {ln}")
    sd = facts["project"].get("seeds", {})
    a(f"- Profile / seeds: {', '.join(f'{k}={v}' for k, v in sd.items())}")
    a(f"- Run time: {facts['project'].get('run_time', '-')}")
    a("")
    path = out / "AUDIT_REPORT.md"
    path.write_text("\n".join(L))
    return path


# ----------------------------------------------------------------------------- docx
def write_docx(facts, checks, env, out: Path, generated: str) -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    SANS, MONO = "Calibri", "JetBrains Mono"
    INK, SOFT, ACC = RGBColor(0x11, 0x18, 0x27), RGBColor(0x4B, 0x55, 0x63), RGBColor(0x08, 0x79, 0x5D)
    p = facts["project"]
    figs = fig_by_id(facts)
    chk = {c["id"]: c for c in checks}

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.orientation = WD_ORIENT.PORTRAIT
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.8))
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    text_w = Cm(21.0 - 3.6)

    st = doc.styles["Normal"]
    st.font.name, st.font.size, st.font.color.rgb = SANS, Pt(10), INK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), SANS)
    st.paragraph_format.space_after = Pt(4)
    for lvl, size in ((1, 18), (2, 14), (3, 12)):
        h = doc.styles[f"Heading {lvl}"]
        h.font.name, h.font.size, h.font.bold = SANS, Pt(size), True
        h.font.color.rgb = INK if lvl != 2 else ACC
        rf = h.element.rPr.rFonts
        for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(att), SANS)
        h.paragraph_format.space_before = Pt(12 if lvl == 1 else 10)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True

    # footer: page X of Y
    def field(run, code):
        for tag, txt in (("begin", None), (None, code), ("end", None)):
            if tag:
                e = OxmlElement("w:fldChar")
                e.set(qn("w:fldCharType"), tag)
            else:
                e = OxmlElement("w:instrText")
                e.set(qn("xml:space"), "preserve")
                e.text = txt
            run._r.append(e)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run(f"{p['name']} · audit report · {env['commit_short']} · page ")
    r.font.size, r.font.color.rgb = Pt(8), SOFT
    r2 = fp.add_run()
    r2.font.size = Pt(8)
    field(r2, "PAGE")
    r3 = fp.add_run(" of ")
    r3.font.size, r3.font.color.rgb = Pt(8), SOFT
    r4 = fp.add_run()
    r4.font.size = Pt(8)
    field(r4, "NUMPAGES")

    def para(text="", bold=False, mono=False, size=None, color=None, italic=False, after=None):
        q = doc.add_paragraph()
        add_runs(q, text, bold=bold, mono=mono, size=size, color=color, italic=italic)
        if after is not None:
            q.paragraph_format.space_after = Pt(after)
        return q

    def add_runs(q, text, bold=False, mono=False, size=None, color=None, italic=False):
        # `code` spans -> monospace
        for i, part in enumerate(re.split(r"`([^`]*)`", str(text))):
            if not part:
                continue
            run = q.add_run(part)
            run.bold, run.italic = bold, italic
            if mono or i % 2 == 1:
                run.font.name = MONO
                run._r.rPr.rFonts.set(qn("w:eastAsia"), MONO)
                run.font.size = Pt((size or 10) - 1.5)
            elif size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color

    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hexcolor)
        tcPr.append(shd)

    def table(headers, rows, widths, mono_cols=(), size=8):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = ""
            add_runs(c.paragraphs[0], h, bold=True, size=size, color=RGBColor(0xFF, 0xFF, 0xFF))
            shade(c, "111827")
        # repeat header row on each page
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trPr.append(th)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                add_runs(cells[i].paragraphs[0], v, mono=i in mono_cols, size=size)
                if ri % 2 == 1:
                    shade(cells[i], "F3F4F6")
        for i, w in enumerate(widths):  # gridCol widths (LibreOffice honours these, not cell widths)
            t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
                for q in row.cells[i].paragraphs:
                    q.paragraph_format.space_after = Pt(0)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    n_pass = sum(c["status"] == "pass" for c in checks)
    n_fail = sum(c["status"] == "FAIL" for c in checks)
    n_man = sum(c["status"] == "manual" for c in checks)

    # ---- cover
    para("AUDIT REPORT", bold=True, size=10, color=ACC, after=2)
    doc.add_heading(p["title"], level=1)
    para(
        "A slide-by-slide record of every number, assumption and calculation in the demo video, with its source.",
        color=SOFT,
        size=11,
    )
    cover = [
        ("Repository", p["repo_url"]),
        (
            "Numbers taken from commit",
            f"`{env['commit']}` (outputs last changed in `{env['outputs_commit']}`)",
        ),
        ("Video", f"`{p['video']}` · {p['video_duration_s']:.1f} s · 1920x1080 · 30 fps · silent"),
        ("Generated", generated),
        ("Automated checks", f"{n_pass} pass · {n_fail} fail · {n_man} manual (of {len(checks)} figures)"),
        ("Assumptions", f"{len(facts['assumptions'])} (A-01 to A-{len(facts['assumptions']):02d})"),
    ]
    table(["Item", "Value"], cover, [4.2, 13.2], size=9)
    doc.add_heading("How to use this report", level=2)
    para(
        "Every number or claim on screen has an ID (F-01, F-02, ...) and every modelling assumption has an ID "
        "(A-01, ...). Cite the ID. Each figure lists its source file, the exact column, row or key, the formula for "
        "derived numbers, and an automated check result. The change-impact matrix (section 5) shows which slides "
        "and files change if an input or assumption changes."
    )
    para(
        "Types: direct = read straight from a repo output; derived = computed from direct values (formula given); "
        "illustrative = drawing or example, not data; general knowledge = context not produced by the repo.",
        color=SOFT,
        size=9,
    )
    doc.add_heading("Regenerate", level=3)
    for c in p.get("regenerate", []):
        para(c, mono=True, size=9, after=1)

    # ---- figures
    doc.add_page_break()
    doc.add_heading("1. Figures", level=2)
    rows = []
    for f in facts["figures"]:
        c = chk[f["id"]]
        status = {"pass": "✓ verified", "FAIL": "✗ FAIL", "manual": "manual"}[c["status"]]
        if c["status"] == "manual":
            status += ": " + c["reason"]
        rows.append(
            [
                f["id"],
                f["shown"],
                ", ".join(f["slides"]),
                f"`{f['source']}`",
                f["key"],
                f["formula"],
                TYPE_LABEL[f["type"]],
                status,
            ]
        )
    table(
        ["ID", "Shown", "Slides", "Source", "Column / row / key", "How computed", "Type", "Check"],
        rows,
        [1.0, 2.6, 1.1, 3.0, 2.9, 2.9, 1.5, 2.4],
        mono_cols=(0,),
        size=7,
    )

    doc.add_heading(f"2. Automated check run ({generated})", level=2)
    rows = [
        [
            c["id"],
            c.get("expect", c["shown"]),
            c["computed"],
            {"pass": "✓ pass", "FAIL": "✗ FAIL", "manual": "manual"}[c["status"]],
        ]
        for c in checks
    ]
    table(
        ["ID", "Expected on slide", "Recomputed from repo", "Result"],
        rows,
        [1.4, 6.4, 6.4, 3.2],
        mono_cols=(0, 1, 2),
        size=7.5,
    )

    doc.add_heading("3. Assumptions register", level=2)
    rows = []
    for x in facts["assumptions"]:
        ids = list(figs) if x.get("all_figures") else x.get("figures", [])
        rows.append(
            [
                x["id"],
                x["statement"],
                f"`{x['source']}`",
                assumption_figs(x, figs),
                f"slides {', '.join(slides_of(ids, figs))}",
            ]
        )
    table(
        ["ID", "Assumption", "Source", "Figures affected", "Slides affected"],
        rows,
        [1.1, 6.2, 4.4, 3.6, 2.1],
        mono_cols=(0,),
        size=7.5,
    )

    # ---- slides
    doc.add_page_break()
    doc.add_heading("4. Slide by slide", level=2)
    for i, s in enumerate(facts["slides"]):
        if i:
            doc.add_page_break()
        doc.add_heading(f"Slide {s['id']}: {s['title']}", level=3)
        para(f"{s['time']} · composition `{s['composition']}`", color=SOFT, size=9)
        for key in ("image", "image2"):
            if s.get(key) and (out / s[key]).exists():
                doc.add_picture(str(out / s[key]), width=text_w)
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        para("On-screen text (verbatim)", bold=True, after=2)
        for t in s["text"]:
            q = para(t, size=9, color=SOFT, after=1)
            q.paragraph_format.left_indent = Cm(0.5)
        on = [f for f in facts["figures"] if s["id"] in f["slides"]]
        if on:
            para("Figures on this slide", bold=True, after=2)
            table(
                ["ID", "Shown", "Source · key", "Calculation", "Check"],
                [
                    [
                        f["id"],
                        f["shown"],
                        f"`{f['source']}` · {f['key']}",
                        f["formula"],
                        {"pass": "✓", "FAIL": "✗", "manual": "manual"}[chk[f["id"]]["status"]],
                    ]
                    for f in on
                ],
                [1.1, 3.4, 6.0, 5.4, 1.5],
                mono_cols=(0,),
                size=7.5,
            )
            for f in on:
                if f.get("note"):
                    para(f"{f['id']} note: {f['note']}", size=8.5, italic=True, color=SOFT)
        para(f"Assumptions: {', '.join(s.get('assumptions', [])) or '-'}", size=9)
        if s.get("illustrative"):
            para("Illustrative, not data:", bold=True, size=9, after=1)
            for t in s["illustrative"]:
                q = para("• " + t, size=9, after=1)
                q.paragraph_format.left_indent = Cm(0.5)
        para(f"To change this slide: {s['change']}", size=9)

    # ---- matrix, caveats, appendix
    doc.add_page_break()
    doc.add_heading("5. Change-impact matrix", level=2)
    table(
        ["Input / assumption", "Affected figures", "Slides", "Files to edit", "Re-run"],
        [
            [f"{r['id']}: {r['input']}", r["figures"], r["slides"], f"`{r['edit']}`", r["rerun"]]
            for r in impact_rows(facts)
        ],
        [4.4, 3.8, 1.8, 4.2, 3.2],
        size=7.5,
    )
    doc.add_heading("6. Known differences and caveats", level=2)
    for c in facts.get("caveats", []):
        q = para("• " + c, size=9.5, after=2)
        q.paragraph_format.left_indent = Cm(0.4)
    doc.add_heading("Appendix: environment", level=2)
    env_rows = [("Commit", f"`{env['commit']}`"), ("Python", env["python"])]
    env_rows += [(ln.split()[0], ln.split()[1]) for ln in env["packages"].splitlines() if ln.strip()]
    sd = facts["project"].get("seeds", {})
    env_rows += [
        ("Profile / seeds", ", ".join(f"{k}={v}" for k, v in sd.items())),
        ("Run time", facts["project"].get("run_time", "-")),
    ]
    if env["dirty"]:
        env_rows.append(("Uncommitted files", env["dirty"].replace("\n", "; ")))
    table(["Item", "Value"], env_rows, [4.2, 13.2], size=8.5)

    path = out / "AUDIT_REPORT.docx"
    doc.save(path)
    return path


def to_pdf(docx: Path) -> Path:
    subprocess.run(
        [SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
        check=True,
        capture_output=True,
    )
    return docx.with_suffix(".pdf")


# ----------------------------------------------------------------------------- main
def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("facts")
    ap.add_argument("--no-stills", action="store_true")
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument(
        "--master", help="rendered master MP4 to take stills from (default: stills.master in the facts file)"
    )
    a = ap.parse_args()
    facts = yaml.safe_load(Path(a.facts).read_text())
    p = facts["project"]
    base = Path(a.facts).resolve().parent  # relative paths in the facts file are relative to it

    def resolve(v: str) -> Path:
        q = Path(v).expanduser()
        return q if q.is_absolute() else (base / q).resolve()

    repo, demo, out = resolve(p["repo_dir"]), resolve(p["demo_dir"]), resolve(p["out_dir"])
    if a.master:
        facts.setdefault("stills", {})["master"] = str(Path(a.master).resolve())
    out.mkdir(parents=True, exist_ok=True)
    generated = dt.datetime.now().strftime("%Y-%m-%d %H:%M")

    ids = [f["id"] for f in facts["figures"]]
    assert len(ids) == len(set(ids)), "duplicate figure ids"
    known = set(ids)
    for x in facts["assumptions"]:
        bad = [i for i in x.get("figures", []) if i not in known]
        assert not bad, f"{x['id']} references unknown figures {bad}"

    if not a.no_stills:
        make_stills(facts, demo, out)
    checks = run_checks(facts, repo, demo)
    env = environment(facts, repo)

    print(f"{'ID':6} {'result':7} {'expected':>28}  {'recomputed':>28}")
    for c in checks:
        print(
            f"{c['id']:6} {c['status']:7} {str(c.get('expect', '-'))[:28]:>28}  {str(c['computed'])[:28]:>28}"
        )
    n_fail = sum(c["status"] == "FAIL" for c in checks)
    print(
        f"\n{sum(c['status'] == 'pass' for c in checks)} pass, {n_fail} fail, "
        f"{sum(c['status'] == 'manual' for c in checks)} manual"
    )

    (out / "audit_checks.json").write_text(
        json.dumps({"generated": generated, "commit": env["commit"], "checks": checks}, indent=1, default=str)
    )
    md = write_md(facts, checks, env, out, generated)
    dx = write_docx(facts, checks, env, out, generated)
    print("wrote", md, dx)
    if not a.no_pdf:
        print("wrote", to_pdf(dx))
    return 1 if n_fail else 0


if __name__ == "__main__":
    sys.exit(main())
